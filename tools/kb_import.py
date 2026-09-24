#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""BizBot 门店知识库导入工具

把店主营维填写好的《门店知识库模板.xlsx》一键转换成 BizBot 可加载的
knowledge_base_data.json（与 config/knowledge_base.py 的加载格式完全一致）。

用法:
    python3 kb_import.py 门店知识库.xlsx
    python3 kb_import.py 门店知识库.xlsx --out /path/to/knowledge_base.json
    python3 kb_import.py 门店知识库.xlsx --include-example   # 连示例行一起导入

模板结构（三张表，列结构一致）:
    FAQ      : 问题 | 触发关键词 | 标准答案 | 排除词 | 推荐追问 | 参考价格
    SOP      : 场景名称 | 触发关键词 | 标准动作/话术 | 排除词 | 推荐追问 | 参考价格
    行业知识 : 主题 | 触发关键词 | 内容说明 | 排除词 | 推荐追问 | 参考价格
    （可选）tuning/调频规则 : 同上结构，优先级最高

规则:
    * 第 1 行为表头，第 2 行为 [示例] 行 —— 默认跳过示例行；
      问题列以「[示例]」开头即视为示例行（--include-example 可保留）。
    * 多个关键词/排除词用逗号（中英文均可）、顿号或分号分隔。
    * 校验：问题/关键词/答案必填，价格填了必须为大于 0 的数字；
      报错会指出工作表名和 Excel 行号，且一次性列出全部问题。
    * 输出 JSON 顶层固定为 {"tuning": [...], "faq": [...], "sop": [...], "industry": [...]},
      条目字段为 question / keywords / answer / follow_up / exclude_keywords，
      与 config/knowledge_base_data.json 现有格式一致，BizBot 加载器无需改动。
"""

import argparse
import json
import os
import re
import sys
from typing import Any, Dict, List, Optional, Tuple

try:
    from openpyxl import load_workbook
except ImportError:  # pragma: no cover
    sys.stderr.write("缺少依赖 openpyxl，请先执行: pip install openpyxl\n")
    sys.exit(2)

# ---------------------------------------------------------------- 常量

# 工作表名 → JSON 顶层 key（匹配时忽略大小写/空白；同一工作表可写中文别名）
SHEET_ALIASES: Dict[str, Tuple[str, ...]] = {
    "faq": ("faq",),
    "sop": ("sop",),
    "industry": ("行业知识", "industry"),
    "tuning": ("tuning", "调频规则", "调频"),
}
CATEGORY_LABEL = {"faq": "FAQ", "sop": "SOP", "industry": "行业知识", "tuning": "调频规则"}
REQUIRED_SHEETS = ("faq", "sop", "industry")
KEY_ORDER = ("faq", "sop", "industry", "tuning")

# 关键词分隔符：中英文逗号 / 顿号 / 分号 / 换行 / 竖线
KEYWORD_SPLIT_RE = re.compile(r"[,，、;；\n|]+")
# 示例行标记（兼容全角方括号）
EXAMPLE_RE = re.compile(r"^\s*[\[［]\s*示\s*例\s*[\]］]")
SEP_DISPLAY = "，/、/；"


def norm_text(value: Any) -> str:
    """单元格值 → 去首尾空白的字符串；None/空 → ''"""
    if value is None:
        return ""
    return str(value).replace("\r\n", "\n").strip()


def split_keywords(raw: str) -> List[str]:
    """按常见分隔符拆关键词，去空、去重（保序）"""
    result: List[str] = []
    for part in KEYWORD_SPLIT_RE.split(raw):
        kw = part.strip()
        if kw and kw not in result:
            result.append(kw)
    return result


def parse_price(raw: str) -> Tuple[Optional[float], Optional[str]]:
    """解析价格列。返回 (价格, 错误信息)；未填返回 (None, None)"""
    if raw == "":
        return None, None
    try:
        price = float(raw)
    except ValueError:
        return None, f"价格「{raw}」不是数字"
    if price <= 0:
        return None, f"价格「{raw}」必须大于 0"
    return price, None


def match_sheet(name: str) -> Optional[str]:
    """Excel 工作表名 → 类别 key；不认识返回 None"""
    key = name.strip().lower()
    for cat, aliases in SHEET_ALIASES.items():
        if key in [a.lower() for a in aliases]:
            return cat
    return None


# ---------------------------------------------------------------- 读取与校验

def read_and_validate(
    path: str, include_example: bool
) -> Tuple[Dict[str, List[Dict[str, Any]]], List[str], List[str]]:
    """读取模板并校验，返回 (数据, errors, warnings)。

    errors 非空时数据不完整，调用方不应写出 JSON。
    """
    errors: List[str] = []
    warnings: List[str] = []
    data: Dict[str, List[Dict[str, Any]]] = {cat: [] for cat in KEY_ORDER}

    if not os.path.exists(path):
        errors.append(f"文件不存在: {path}")
        return data, errors, warnings

    try:
        wb = load_workbook(path, read_only=True, data_only=True)
    except Exception as e:  # noqa: BLE001
        errors.append(f"无法打开 Excel 文件（请确认是 .xlsx 模板）: {e}")
        return data, errors, warnings

    seen_sheets: set = set()
    for ws in wb.worksheets:
        cat = match_sheet(ws.title)
        if cat is None:
            warnings.append(f"工作表「{ws.title}」无法识别，已跳过（模板标准表：FAQ / SOP / 行业知识）")
            continue
        seen_sheets.add(cat)

        rows = list(ws.iter_rows(min_row=1, max_col=6, values_only=True))
        if not rows:
            continue
        # 第 1 行视为表头（不校验内容，允许店主管改表头文字）
        seen_questions: set = set()
        for row_idx, row in enumerate(rows[1:], start=2):  # Excel 实际行号
            question_raw = norm_text(row[0])
            keywords_raw = norm_text(row[1])
            answer_raw = norm_text(row[2])
            exclude_raw = norm_text(row[3])
            follow_up_raw = norm_text(row[4])
            price_raw = norm_text(row[5])

            # 整行全空 → 静默跳过
            if not any((question_raw, keywords_raw, answer_raw, exclude_raw, follow_up_raw, price_raw)):
                continue

            where = f"[{CATEGORY_LABEL[cat]}] 第 {row_idx} 行"
            is_example = bool(EXAMPLE_RE.match(question_raw))
            if is_example and not include_example:
                continue

            # ---- 必填校验 ----
            row_errors: List[str] = []
            if not question_raw:
                row_errors.append("问题/场景 不能为空")
            if not keywords_raw:
                row_errors.append("触发关键词 不能为空")
            if not answer_raw:
                row_errors.append("标准答案/内容 不能为空")

            keywords = split_keywords(keywords_raw)
            if keywords_raw and not keywords:
                row_errors.append(f"触发关键词没有可用的词（请用 {SEP_DISPLAY} 分隔）")

            # ---- 价格校验 ----
            price, price_err = parse_price(price_raw)
            if price_err:
                row_errors.append(price_err)
            if price is not None and not re.search(r"\d", answer_raw):
                warnings.append(
                    f"{where}: 填了参考价格但答案里没写数字，机器人报不出价格，"
                    f"建议把具体价格写进\"标准答案/内容\"列"
                )

            for msg in row_errors:
                errors.append(f"{where}: {msg}")
            if row_errors:
                continue

            # ---- 重复提示（warning，不阻断）----
            q_key = question_raw.lower()
            if q_key in seen_questions:
                warnings.append(f"{where}: 问题「{question_raw}」重复出现，建议合并或修改")
            seen_questions.add(q_key)

            # ---- 组装条目（字段与现有 knowledge_base_data.json 一致）----
            item: Dict[str, Any] = {"question": question_raw, "keywords": keywords}
            exclude = split_keywords(exclude_raw)
            if exclude:
                item["exclude_keywords"] = exclude
            item["answer"] = answer_raw
            if follow_up_raw:
                item["follow_up"] = follow_up_raw
            data[cat].append(item)

    wb.close()

    missing = [CATEGORY_LABEL[c] for c in REQUIRED_SHEETS if c not in seen_sheets]
    if missing:
        errors.append(f"缺少标准工作表: {'、'.join(missing)}（请使用配套的《门店知识库模板.xlsx》）")

    if not errors and not any(data.values()):
        warnings.append("没有读到任何数据行（只有表头），请检查表格是否填写")
    return data, errors, warnings


# ---------------------------------------------------------------- 输出

def build_json(data: Dict[str, List[Dict[str, Any]]]) -> Dict[str, List[Dict[str, Any]]]:
    """组装顶层 JSON，key 顺序与 BizBot _save_to_file 一致"""
    return {cat: data.get(cat, []) for cat in ("tuning", "faq", "sop", "industry")}


def main() -> int:
    parser = argparse.ArgumentParser(
        description="把《门店知识库模板.xlsx》转换成 BizBot 知识库 JSON",
        formatter_class=argparse.RawDescriptionHelpFormatter,
        epilog="示例: python3 kb_import.py 门店知识库.xlsx --out knowledge_base.json",
    )
    parser.add_argument("input", help="填写好的门店知识库 Excel 文件（.xlsx）")
    parser.add_argument("--out", default="knowledge_base.json", help="输出 JSON 路径（默认 knowledge_base.json）")
    parser.add_argument(
        "--include-example", action="store_true", help="保留 [示例] 行（默认跳过）"
    )
    args = parser.parse_args()

    data, errors, warnings = read_and_validate(args.input, args.include_example)

    for w in warnings:
        print(f"⚠ 提示: {w}")

    if errors:
        print(f"\n✗ 发现 {len(errors)} 个问题，请修改后重新导入:")
        for i, e in enumerate(errors, 1):
            print(f"  {i}. {e}")
        return 1

    payload = build_json(data)
    out_path = os.path.abspath(args.out)
    out_dir = os.path.dirname(out_path)
    if out_dir:
        os.makedirs(out_dir, exist_ok=True)
    with open(out_path, "w", encoding="utf-8") as f:
        json.dump(payload, f, ensure_ascii=False, indent=2)

    total = 0
    print("\n✓ 导入完成:")
    for cat in KEY_ORDER:
        n = len(payload[cat])
        total += n
        print(f"  {CATEGORY_LABEL[cat]:<6} {n} 条")
    print(f"  合计   {total} 条")
    print(f"输出文件: {out_path}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
